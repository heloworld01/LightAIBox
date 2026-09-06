"""沙箱内安全的文件产出工具：让智能体把成果真正写成文件。

与 DesktopTools（app/agent_tools.py）同构：每个工具是 dict ToolDef（name / description /
parameters(JSON Schema) / run(dict)->str），经 gateway_llm 的 DesktopToolAdapter 接入
LightAgents ToolRegistry，从而自动出现在 schema 与「工具调用步骤」展示中。

安全设计（对应用户已确认的「结构化生成 + 每次确认」决策）：
- **结构化生成**：write_text / write_docx / write_xlsx 走 python-docx / openpyxl 直接生成，
  不执行任意脚本、不依赖外部 CLI。
- **沙箱**：所有路径解析严格限制在本会话目录内（拒绝对路径 / `..` / 越界，commonpath 校验），
  扩展名受限（docx / xlsx 强制），单文件大小上限 max_bytes。
- **每次确认**：写盘前调 approval.await_approval（app/approval.py），由 UI 主线程弹窗，
  用户拒绝则不落任何文件。
- **无副作用**：工具只在其沙箱目录内生成文件。
"""
import os
from typing import Callable, Dict, List

# 每个工具：name / description / parameters(JSON Schema) / run(dict)->str
ToolDef = Dict[str, object]


def _human_size(n: int) -> str:
    if n < 1024:
        return f"{n} 字节"
    if n < 1024 * 1024:
        return f"{n / 1024:.1f} KB"
    return f"{n / 1024 / 1024:.1f} MB"


class FileTools:
    """会话沙箱内的文件写出工具集合。session_dir 与 approval 必须同时提供才能注册。"""

    def __init__(self, session_dir: str, approval, max_bytes: int = 20 * 1024 * 1024):
        self.session_dir = os.path.realpath(session_dir)
        self.approval = approval
        self.max_bytes = max_bytes
        self._tools: Dict[str, ToolDef] = {}
        self._register("write_text", "写入纯文本 / Markdown 到沙箱文件", {
            "type": "object",
            "properties": {
                "path": {"type": "string",
                         "description": "会话目录内的相对路径，如 'report.md' 或 'out/notes.txt'"},
                "content": {"type": "string", "description": "要写入的文本内容"},
            },
            "required": ["path", "content"],
        }, self._write_text)
        self._register("write_docx", "生成一份 Word (.docx) 文档：标题 + 若干段落", {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "相对路径，必须以 .docx 结尾，如 'report.docx'"},
                "title": {"type": "string", "description": "文档标题（可空）"},
                "paragraphs": {"type": "array", "items": {"type": "string"},
                               "description": "正文段落列表"},
            },
            "required": ["path"],
        }, self._write_docx)
        self._register("write_xlsx", "生成一份 Excel (.xlsx) 表格：首行作表头", {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "相对路径，必须以 .xlsx 结尾，如 'data.xlsx'"},
                "sheet_name": {"type": "string", "description": "工作表名，默认 Sheet1"},
                "rows": {"type": "array", "items": {"type": "array",
                                                    "description": "一行，元素为字符串或数字"},
                         "description": "二维数组，首行作为表头"},
            },
            "required": ["path"],
        }, self._write_xlsx)

    # ------------------------------------------------------------------ #
    def _register(self, name: str, description: str, parameters: dict,
                  run: Callable[[dict], str]) -> None:
        self._tools[name] = {
            "name": name,
            "description": description,
            "parameters": parameters,
            "run": run,
        }

    def list_tools(self) -> List[ToolDef]:
        return list(self._tools.values())

    def schemas_openai(self) -> List[dict]:
        return [
            {"type": "function",
             "function": {"name": t["name"], "description": t["description"],
                          "parameters": t["parameters"]}}
            for t in self._tools.values()
        ]

    def schemas_anthropic(self) -> List[dict]:
        return [
            {"name": t["name"], "description": t["description"],
             "input_schema": t["parameters"]}
            for t in self._tools.values()
        ]

    def execute(self, name: str, args: dict) -> str:
        tool = self._tools.get(name)
        if tool is None:
            return f"❌ 未找到工具 '{name}'"
        run: Callable[[dict], str] = tool["run"]  # type: ignore[assignment]
        try:
            return run(args or {})
        except ValueError as exc:  # 沙箱/扩展名校验失败 → 直接作为观察结果
            return f"❌ {exc}"
        except ImportError as exc:
            return f"❌ 缺少依赖生成该格式（python-docx / openpyxl）：{exc}"
        except Exception as exc:  # noqa: BLE001 —— 工具异常转成观察结果喂回模型
            return f"❌ 工具执行失败：{exc}"

    # ------------------------------------------------------------------ #
    def ensure_dirs(self) -> None:
        os.makedirs(self.session_dir, exist_ok=True)

    def resolve(self, raw: str) -> str:
        """把相对路径解析为沙箱内的绝对路径；越界/绝对路径/`..` 一律拒绝。

        返回 realpath 后的绝对路径，保证位于 self.session_dir 之下。
        """
        if not raw or not isinstance(raw, str):
            raise ValueError("路径无效")
        raw = raw.strip().replace("\\", "/")
        if not raw:
            raise ValueError("路径为空")
        if os.path.isabs(raw) or raw.startswith("/") or raw.startswith("~") or \
                (len(raw) >= 2 and raw[1] == ":"):
            raise ValueError("只允许会话目录内的相对路径")
        parts = [p for p in raw.split("/") if p not in ("", ".")]
        if ".." in parts:
            raise ValueError("路径不允许越界（..）")
        root = self.session_dir
        target = os.path.realpath(os.path.join(root, *parts))
        # realpath 后的 target 必须仍在沙箱根下（可防符号链接逃逸 / 硬到其他位置）
        if os.path.commonpath([root, target]) != root:
            raise ValueError("路径越出会话目录，已拒绝")
        return target

    def _commit(self, kind: str, action: str, path: str, data: bytes) -> str:
        """构造授权提案（路径/类型/大小）→ await_approval → 通过则写盘。"""
        size = len(data)
        if size > self.max_bytes:
            return f"❌ 文件拟写入 {self._human(size)}，超过单文件上限 {self._human(self.max_bytes)}，已拒绝"
        proposal = {
            "action": action,
            "kind": kind,
            "path": path,
            "display": self._short(path),
            "size": size,
        }
        ok = self.approval.await_approval(proposal)
        if not ok:
            return f"❌ 用户拒绝了写入：{self._short(path)}"
        self.ensure_dirs()
        os.makedirs(os.path.dirname(path) or self.session_dir, exist_ok=True)
        try:
            with open(path, "wb") as f:
                f.write(data)
        except Exception as exc:  # noqa: BLE001
            return f"❌ 写入失败：{exc}"
        return f"已写入：{path}（{self._human(size)}）"

    @staticmethod
    def _human(n: int) -> str:
        return _human_size(n)

    @staticmethod
    def _short(path: str) -> str:
        return os.path.basename(path)

    # ------------------------------------------------------------------ #
    def _write_text(self, args: dict) -> str:
        path = self.resolve(args.get("path", ""))
        content = args.get("content", "")
        if isinstance(content, (list, tuple)):
            content = "\n".join(str(x) for x in content)
        data = str(content or "").encode("utf-8")
        return self._commit("txt", "写入文本文件", path, data)

    def _write_docx(self, args: dict) -> str:
        path = self.resolve(args.get("path", ""))
        if not path.lower().endswith(".docx"):
            raise ValueError("write_docx 仅支持 .docx 扩展名（路径需以 .docx 结尾）")
        import io
        from docx import Document  # python-docx

        title = str(args.get("title", "") or "")
        paragraphs = args.get("paragraphs") or []
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]
        doc = Document()
        if title:
            doc.add_heading(title, level=0)
        for p in paragraphs:
            doc.add_paragraph(str(p))
        bio = io.BytesIO()
        doc.save(bio)
        return self._commit("docx", "生成 Word 文档", path, bio.getvalue())

    def _write_xlsx(self, args: dict) -> str:
        path = self.resolve(args.get("path", ""))
        if not path.lower().endswith(".xlsx"):
            raise ValueError("write_xlsx 仅支持 .xlsx 扩展名（路径需以 .xlsx 结尾）")
        import io
        from openpyxl import Workbook

        sheet_name = str(args.get("sheet_name") or "Sheet1")[:31]
        rows = args.get("rows") or []
        if isinstance(rows, dict):  # 兼容 {rows:[...]} 包装
            rows = rows.get("rows") or rows.get("data") or []
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in rows:
            ws.append(list(row) if isinstance(row, (list, tuple)) else [row])
        bio = io.BytesIO()
        wb.save(bio)
        return self._commit("xlsx", "生成 Excel 表格", path, bio.getvalue())
